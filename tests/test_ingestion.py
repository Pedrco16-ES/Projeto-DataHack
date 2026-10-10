"""Fixtures didáticas restritas aos testes; não são dados oficiais."""

from io import BytesIO
import json
from pathlib import Path
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from src.export.administrative import csv_bytes, xlsx_bytes
from src.ingestion.pipeline import process_file
from src.ingestion.readers import preview
from src.ingestion.schema import IngestionError, load_contract
from src.processing.validation import clean_strings, validate_rows


def census(**changes):
    row = dict(NU_ANO_CENSO="2024", CO_IES="0001", CO_CURSO="00123", CO_UF="51",
               CO_MUNICIPIO="5103403", TP_DIMENSAO="1", TP_MODALIDADE_ENSINO="1",
               TP_CATEGORIA_ADMINISTRATIVA="1", QT_MAT="30", QT_ING="20", QT_CONC="10", QT_VG_TOTAL="40")
    return {**row, **changes}


def trajectory(**changes):
    row = dict(CO_IES="0001", CO_CURSO="00123", CO_UF="51", CO_MUNICIPIO="5103403",
               NU_ANO_INGRESSO="2015", NU_ANO_REFERENCIA="2024", TP_MODALIDADE_ENSINO="1",
               QT_INGRESSANTE="100", QT_PERMANENCIA="10", QT_CONCLUINTE="20", QT_DESISTENCIA="30",
               QT_FALECIDO="0", TAP="10", TCA="40", TDA="50", TCAN="20", TADA="30")
    return {**row, **changes}


def save_csv(tmp_path, rows, name="source.csv"):
    path = tmp_path / name
    pd.DataFrame(rows).to_csv(path, sep=";", index=False, encoding="latin-1")
    return path


def run(tmp_path, path, **kwargs):
    return process_file(path, "censo_curso", 2024, output_root=tmp_path / "processed", **kwargs)


def table(report, name="accepted.csv"):
    return pd.read_csv(Path(report["batch_dir"]) / name, dtype="string", keep_default_na=False)


def test_csv_mt_codes_missing_and_reconciliation(tmp_path):
    path = save_csv(tmp_path, [census(QT_VG_TOTAL="."), census(CO_UF="35"),
                               census(CO_UF="", CO_MUNICIPIO="", TP_DIMENSAO="3", TP_MODALIDADE_ENSINO="2")])
    report = run(tmp_path, path, chunksize=1)
    assert report["counts"] == dict(received=3, accepted=1, rejected=0, excluded=2, mt_candidates=1, duplicate_rows=0)
    assert table(report).iloc[0].CO_CURSO == "00123"
    assert table(report).iloc[0].QT_VG_TOTAL == ""
    assert report["quality"]["missing"]["QT_VG_TOTAL"] == 1
    assert report["database_rows_imported"] == 0


def test_duplicates_across_chunks_quarantine_all_and_keep_municipal_grain(tmp_path):
    path = save_csv(tmp_path, [census(), census(CO_MUNICIPIO="5108402"), census()])
    result = run(tmp_path, path, chunksize=1)
    assert result["counts"]["accepted"] == 1
    assert result["counts"]["duplicate_rows"] == 2
    assert result["counts"]["rejected"] == 2
    assert table(result).iloc[0].CO_MUNICIPIO == "5108402"


def test_duplicate_invalid_record_also_blocks_valid_occurrence(tmp_path):
    path = save_csv(tmp_path, [census(), census(QT_MAT="-1")])
    result = run(tmp_path, path, chunksize=1)
    assert result["counts"]["accepted"] == 0
    assert result["counts"]["rejected"] == 2


@pytest.mark.parametrize("change,reason", [
    ({"QT_MAT": "-1"}, "fora_da_faixa"), ({"QT_MAT": "3.5"}, "inteiro_invalido"),
    ({"CO_CURSO": ""}, "chave_ausente"), ({"CO_CURSO": "ABC"}, "codigo_invalido"),
    ({"NU_ANO_CENSO": "2023"}, "ano_incompativel"), ({"CO_MUNICIPIO": "3500000"}, "municipio_incompativel"),
    ({"CO_UF": "XX"}, "UF_ausente_ou_invalida"), ({"QT_ING": "inválido"}, "numero_invalido"),
])
def test_rejections_preserve_values_and_reason(tmp_path, change, reason):
    report = run(tmp_path, save_csv(tmp_path, [census(**change)]))
    assert report["counts"]["rejected"] == 1
    assert reason in table(report, "rejected.csv").iloc[0]["_reasons"]
    for key, value in change.items():
        assert table(report, "rejected.csv").iloc[0][key] == value


def test_repeated_file_reuses_verified_batch_and_tampering_blocks(tmp_path):
    path = save_csv(tmp_path, [census()])
    first = run(tmp_path, path, chunksize=1)
    second = run(tmp_path, path, chunksize=10)
    assert second["reused"] and first["batch_id"] == second["batch_id"]
    assert first["artifacts"] == second["artifacts"]
    (Path(first["batch_dir"]) / "accepted.csv").write_text("alterado", encoding="utf-8")
    with pytest.raises(IngestionError, match="alterado"):
        run(tmp_path, path)


def test_xlsx_header_ninth_row_and_all_reference_years(tmp_path):
    path = tmp_path / "source.xlsx"
    book = Workbook(); sheet = book.active
    for _ in range(8):
        sheet.append(["Título didático"])
    rows = [trajectory(NU_ANO_REFERENCIA="2015"), trajectory()]
    sheet.append(list(rows[0])); [sheet.append(list(r.values())) for r in rows]
    book.save(path); book.close()
    frame, layout = preview(path, load_contract("trajetoria", 2024))
    assert layout["header"] == 8 and len(frame) == 2
    result = process_file(path, "trajetoria", 2024, output_root=tmp_path / "processed", chunksize=1)
    assert result["counts"]["accepted"] == 2
    assert table(result).ANOS_DESDE_INGRESSO.tolist() == ["0", "9"]
    assert table(result).CO_CURSO.tolist() == ["00123", "00123"]


def test_trajectory_rate_and_period_validation():
    contract = load_contract("trajetoria", 2024)
    frame = clean_strings(pd.DataFrame([trajectory(TDA="101"), trajectory(NU_ANO_REFERENCIA="2014"), trajectory(TAP="10,5")]))
    cleaned, reasons = validate_rows(frame, contract)
    assert "fora_da_faixa:TDA" in reasons.iloc[0]
    assert "periodo_incompativel" in reasons.iloc[1]
    assert cleaned.TAP.iloc[2] == 10.5


def test_wrong_source_and_missing_column_rejected_before_publication(tmp_path):
    row = census(); del row["CO_CURSO"]
    with pytest.raises(IngestionError, match="Cabeçalho"):
        run(tmp_path, save_csv(tmp_path, [row]))
    assert not list((tmp_path / "processed").glob("*/report.json"))


def test_manual_mapping_requires_recorded_confirmation(tmp_path):
    row = census(); row["CURSO_ORIGINAL"] = row.pop("CO_CURSO")
    path = save_csv(tmp_path, [row])
    mapping = {"CURSO_ORIGINAL": "CO_CURSO"}
    with pytest.raises(IngestionError, match="confirmação"):
        run(tmp_path, path, mapping=mapping)
    report = run(tmp_path, path, mapping=mapping, confirm_mapping=True, mapping_reason="Fixture confirmada pelo teste")
    assert report["mapping"] == mapping
    assert table(report).CO_CURSO.iloc[0] == "00123"


def test_json_preserves_identifier_strings(tmp_path):
    path = tmp_path / "source.json"
    path.write_text(json.dumps([census()]), encoding="utf-8")
    assert table(run(tmp_path, path)).CO_CURSO.iloc[0] == "00123"


def test_zip_single_member_and_path_traversal(tmp_path):
    original = save_csv(tmp_path, [census()])
    path = tmp_path / "source.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.write(original, "dados/source.csv")
    assert run(tmp_path, path)["counts"]["accepted"] == 1
    unsafe = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(unsafe, "w") as z:
        z.write(original, "../source.csv")
    with pytest.raises(IngestionError, match="inseguro"):
        run(tmp_path, unsafe)


def test_exports_preserve_text_ids_and_neutralize_formulas():
    frame = pd.DataFrame({"codigo": ["00123"], "texto": ["=1+1"]})
    assert "'=1+1" in csv_bytes(frame).decode("utf-8-sig")
    book = load_workbook(BytesIO(xlsx_bytes(frame)), read_only=True)
    row = list(book.active.values)[1]
    assert row == ("00123", "'=1+1")
    book.close()


def test_contracts_are_specific_to_edition():
    assert len(load_contract("censo_curso", 2021)["columns"]) == 200
    assert len(load_contract("censo_curso", 2024)["columns"]) == 223
    with pytest.raises(IngestionError, match="Edição"):
        load_contract("censo_curso", 2020)


def test_numeric_excel_codes_with_nulls_do_not_become_decimal_strings(tmp_path):
    path = tmp_path / "numeric.xlsx"
    first = trajectory(CO_UF=51, CO_CURSO=123, CO_IES=1, CO_MUNICIPIO=5103403)
    second = trajectory(CO_UF=None, CO_CURSO=None, CO_IES=None)
    book = Workbook(); sheet = book.active
    sheet.append(list(first)); sheet.append(list(first.values())); sheet.append(list(second.values()))
    book.save(path); book.close()
    report = process_file(path, "trajetoria", 2024, output_root=tmp_path / "processed", chunksize=2)
    assert report["counts"]["accepted"] == 1
    assert table(report).CO_CURSO.iloc[0] == "123"
    other = process_file(path, "trajetoria", 2024, output_root=tmp_path / "another", chunksize=1)
    assert other["artifacts"] == report["artifacts"]


def test_national_dimension_with_zero_uf_is_not_assigned_to_mt(tmp_path):
    path = save_csv(tmp_path, [census(CO_UF="0", CO_MUNICIPIO="0", TP_DIMENSAO="3", TP_MODALIDADE_ENSINO="2")])
    report = run(tmp_path, path)
    assert report["counts"]["excluded"] == 1
    assert report["excluded_reasons"]["sem_alocacao_geografica"] == 1


def test_csv_extra_field_fails_instead_of_shifting_columns(tmp_path):
    path = save_csv(tmp_path, [census()])
    text = path.read_text(encoding="latin-1").splitlines()
    path.write_text(text[0] + "\n" + text[1] + ";unexpected\n", encoding="latin-1")
    with pytest.raises(IngestionError):
        run(tmp_path, path)
