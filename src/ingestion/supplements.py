"""Preparação privada das novas fontes. Não modifica o snapshot do dashboard."""

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.analytics.build import normalize
from src.analytics.metrics import AnalyticsError, trajectory, trajectory_summary
from src.config import PROJECT_ROOT
from src.ingestion.pipeline import sha256
from src.ingestion.schema import load_contract
from src.processing.validation import clean_strings, validate_rows

RAW = PROJECT_ROOT / "data/raw/complementares_20261010"
DEST = PROJECT_ROOT / "data/processed/complementares_20261010"


def checked_source(name):
    manifest = json.loads((RAW / "manifest.json").read_text(encoding="utf-8"))
    path = RAW / name
    if name not in manifest["files"] or not path.exists() or sha256(path) != manifest["files"][name]["sha256"]:
        raise AnalyticsError("Fonte ausente ou hash divergente: " + name)
    return path


def population_from_json(data):
    if len(data) != 1 or data[0]["id"] != "93" or data[0]["unidade"] != "Pessoas":
        raise AnalyticsError("Variável populacional inesperada.")
    rows = []
    expected_ages = {str(6575 + i): 18 + i for i in range(7)}
    seen_ages = set()
    for result in data[0]["resultados"]:
        classifications = {item["id"]: item["categoria"] for item in result["classificacoes"]}
        if classifications.get("2") != {"6794": "Total"} or classifications.get("286") != {"113635": "Total"}:
            raise AnalyticsError("Sexo/declaração de idade devem ser Total.")
        ages = classifications.get("287", {})
        if len(ages) != 1:
            raise AnalyticsError("Idade simples esperada.")
        age_id = next(iter(ages))
        if age_id not in expected_ages or age_id in seen_ages:
            raise AnalyticsError("Idade inesperada ou repetida.")
        age = expected_ages[age_id]
        if ages[age_id] != f"{age} anos":
            raise AnalyticsError("Rótulo etário incompatível.")
        seen_ages.add(age_id)
        for item in result["series"]:
            locality = item["localidade"]
            if locality["nivel"]["id"] != "N6" or not re.fullmatch(r"51\d{5}", locality["id"]):
                raise AnalyticsError("Município fora de MT.")
            if set(item["serie"]) != {"2022"}:
                raise AnalyticsError("População deve ser de 2022.")
            value = item["serie"]["2022"]
            # SIDRA '-' significa zero absoluto. '..', '...' e 'X' são ausências,
            # indisponibilidade/sigilo e não são convertidos em zero.
            if value == "-":
                value = "0"
            if not re.fullmatch(r"\d+", value):
                raise AnalyticsError("População ausente, sigilosa ou não numérica.")
            rows.append({"codigo_municipio": locality["id"], "municipio": locality["nome"].removesuffix(" - MT"),
                         "ano": 2022, "idade": age, "populacao": int(value)})
    detail = pd.DataFrame(rows)
    if seen_ages != set(expected_ages) or detail.duplicated(["codigo_municipio", "idade"]).any():
        raise AnalyticsError("Idades incompletas ou chaves repetidas.")
    if detail.groupby("codigo_municipio").idade.nunique().ne(7).any():
        raise AnalyticsError("Município sem as sete idades.")
    grouped = detail.groupby(["codigo_municipio", "municipio", "ano"], as_index=False).populacao.sum()
    return detail, grouped.rename(columns={"populacao": "populacao_18_24"})


def validate_geometry(geometry, codes):
    if geometry.get("type") != "FeatureCollection":
        raise AnalyticsError("Geometria não é FeatureCollection.")
    found = [f["properties"]["codarea"] for f in geometry["features"]]
    if len(found) != len(set(found)) or set(found) != set(codes):
        raise AnalyticsError("Malha e municípios do Censo 2022 não correspondem.")
    def rings(coords, kind):
        if kind == "Polygon":
            return coords
        if kind == "MultiPolygon":
            return [ring for polygon in coords for ring in polygon]
        raise AnalyticsError("Geometria não poligonal.")
    for feature in geometry["features"]:
        for ring in rings(feature["geometry"]["coordinates"], feature["geometry"]["type"]):
            if len(ring) < 4 or ring[0] != ring[-1]:
                raise AnalyticsError("Anel geográfico aberto ou inválido.")
            if any(len(p) < 2 or not (-62 <= p[0] <= -50 and -19 <= p[1] <= -7) for p in ring):
                raise AnalyticsError("Coordenadas incompatíveis com MT em longitude/latitude.")


def enade_from_frame(frame):
    frame = frame.copy()
    frame.columns = [normalize(c).rstrip("0123456789 ") for c in frame]
    mapping = {
        "ano": "ano", "codigo da ies": "CO_IES", "codigo do curso": "CO_CURSO",
        "codigo da area": "codigo_area", "area de avaliacao": "area", "grau academico": "grau",
        "nome da ies": "instituicao", "categoria administrativa": "categoria",
        "modalidade de ensino": "modalidade", "codigo do municipio": "codigo_municipio",
        "municipio do curso": "municipio", "sigla da uf": "uf",
        "no de concluintes inscritos": "inscritos", "no de concluintes participantes": "participantes",
        "total de concluinte igual ou acima do padrao 1 de proficiencia": "proficientes",
        "percentual de concluintes igual ou acima do padrao 1 de proficiencia": "fracao_oficial",
        "conceito enade (faixa)": "conceito",
    }
    if not set(mapping).issubset(frame):
        raise AnalyticsError("Layout Enade inesperado.")
    # A planilha tem notas de rodapé após os registros: ano e códigos vazios.
    year = pd.to_numeric(frame.ano, errors="coerce")
    if (year.isna() & frame["codigo do curso"].notna()).any():
        raise AnalyticsError("Curso com ano ausente.")
    rows = frame.loc[year.notna(), list(mapping)].rename(columns=mapping).copy()
    for c in ["CO_IES", "CO_CURSO", "codigo_area", "codigo_municipio"]:
        rows[c] = rows[c].astype("string").str.strip().str.replace(r"\.0$", "", regex=True)
        if rows[c].isna().any() or not rows[c].str.fullmatch(r"\d+").all():
            raise AnalyticsError("Código inválido no Enade: " + c)
    if rows.duplicated(["ano", "CO_IES", "CO_CURSO"]).any():
        raise AnalyticsError("Chaves repetidas no Enade.")
    for c in ["ano", "inscritos", "participantes", "proficientes", "fracao_oficial"]:
        original = rows[c].astype("string").str.strip().str.replace(",", ".", regex=False)
        rows[c] = pd.to_numeric(original, errors="coerce")
        if (original.notna() & rows[c].isna()).any() or rows[c].dropna().lt(0).any():
            raise AnalyticsError("Valor numérico inválido no Enade.")
    if not rows.ano.eq(2025).all() or rows.inscritos.isna().any():
        raise AnalyticsError("Ano/inscritos inválidos no Enade.")
    for c in ["inscritos", "participantes", "proficientes"]:
        if rows[c].dropna().mod(1).ne(0).any():
            raise AnalyticsError("Contagem fracionária no Enade.")
    if (rows.participantes > rows.inscritos).any() or (rows.proficientes > rows.participantes).any():
        raise AnalyticsError("Numeradores Enade excedem denominadores.")
    complete = rows[["participantes", "proficientes", "fracao_oficial"]].notna().all(axis=1)
    if rows.loc[complete, "participantes"].le(0).any():
        raise AnalyticsError("Denominador Enade não positivo.")
    calc = rows.loc[complete, "proficientes"] / rows.loc[complete, "participantes"]
    # Fonte armazena FRAÇÃO, arredondada a três casas, apesar do título Percentual.
    if not np.allclose(calc.astype(float), rows.loc[complete, "fracao_oficial"].astype(float), atol=0.00051, rtol=0):
        raise AnalyticsError("Fração publicada diverge de proficientes/participantes.")
    incomplete = ~complete
    if not rows.loc[incomplete, "conceito"].eq("SC").all():
        raise AnalyticsError("Resultado incompleto sem indicação SC.")
    rows["proficiencia_pct"] = rows.proficientes / rows.participantes.where(rows.participantes.gt(0)) * 100
    rows["status"] = np.where(complete, "calculavel", "sem_conceito_ou_dado")
    mt = rows.loc[rows.uf.eq("MT")].copy()
    if not mt.codigo_municipio.str.startswith("51").all():
        raise AnalyticsError("UF/município Enade incompatíveis.")
    return mt, {"cursos_brasil": len(rows), "rodapes_excluidos": len(frame) - len(rows), "cursos_mt": len(mt),
                "sem_resultado_mt": int(mt.status.ne("calculavel").sum()),
                "modalidades_mt": {str(k): int(v) for k, v in mt.modalidade.value_counts().items()}}


def prepare_ibge():
    detail, population = population_from_json(json.loads(checked_source("ibge_18_24_municipios_mt_2022.json").read_text(encoding="utf-8")))
    geometry = json.loads(checked_source("ibge_malha_mt_2022.geojson").read_text(encoding="utf-8"))
    validate_geometry(geometry, population.codigo_municipio)
    if len(population) != 141 or population.populacao_18_24.le(0).any():
        raise AnalyticsError("Cadastro 2022 incompleto ou população não positiva.")
    detail.to_csv(DEST / "ibge_idades_mt_2022.csv", index=False)
    population.to_csv(DEST / "ibge_populacao_mt_2022.csv", index=False)
    population[["codigo_municipio", "municipio", "ano"]].to_csv(DEST / "municipios_mt_2022.csv", index=False)
    (DEST / "municipios_mt_2022.geojson").write_text(json.dumps(geometry, ensure_ascii=False), encoding="utf-8")
    current = json.loads(checked_source("ibge_municipios_mt_atual.json").read_text(encoding="utf-8"))
    return {"municipios": len(population), "idades_por_municipio": 7, "linhas_idade": len(detail),
            "populacao_18_24": int(population.populacao_18_24.sum()), "geometrias": len(geometry["features"]),
            "municipios_atuais_fora_2022": [{"codigo": str(x["id"]), "nome": x["nome"]} for x in current if str(x["id"]) not in set(population.codigo_municipio)]}


def prepare_enade():
    frame = pd.read_excel(checked_source("enade_licenciaturas_2025.xlsx"), sheet_name="Conceito Enade Licenciaturas", dtype="string")
    mt, report = enade_from_frame(frame)
    mt.to_csv(DEST / "enade_licenciaturas_mt_2025.csv", index=False)
    return report


def prepare_cohort(year):
    source = checked_source(f"trajetoria_{year}_2024.zip")
    output = DEST / f"trajetoria_mt_{year}_2024.csv"
    info_path = DEST / f"trajetoria_mt_{year}_2024.report.json"
    parser_hash = sha256(Path(__file__))
    if output.exists() and info_path.exists():
        previous = json.loads(info_path.read_text(encoding="utf-8"))
        if previous.get("source_sha256") == sha256(source) and previous.get("output_sha256") == sha256(output) and previous.get("parser_sha256") == parser_hash:
            print(f"Coorte {year}: recorte validado reutilizado", flush=True)
            return previous
    contract = load_contract("trajetoria", 2024)
    rows, total, unknown = [], 0, 0
    with ZipFile(source) as archive:
        names = [n for n in archive.namelist() if n.lower().endswith(".xlsx") and not Path(n).name.startswith("~$")]
        dictionaries = [n for n in archive.namelist() if n.lower().endswith(".docx")]
        if len(names) != 1 or len(dictionaries) != 1:
            raise AnalyticsError("ZIP de coorte com layout ambíguo.")
        dictionary_hash = hashlib.sha256(archive.read(dictionaries[0])).hexdigest()
        if dictionary_hash != contract["dictionary_sha256"]:
            raise AnalyticsError("Dicionário da nova coorte difere do contrato; revisão necessária.")
        with archive.open(names[0]) as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            try:
                sheet = book["INDICADORES_TRAJETORIA"]
                iterator = sheet.iter_rows(min_row=9, values_only=True)
                header = list(next(iterator))
                if header != list(contract["columns"]):
                    raise AnalyticsError("Colunas da nova coorte diferem do contrato.")
                uf_index = header.index("CO_UF")
                for row in iterator:
                    if all(v is None for v in row):
                        continue
                    total += 1
                    if row[uf_index] in [51, "51"]:
                        rows.append(row)
                    elif row[uf_index] is None:
                        unknown += 1
                    if total % 100000 == 0:
                        print(f"Coorte {year}: {total} registros lidos", flush=True)
            finally:
                book.close()
    frame = clean_strings(pd.DataFrame(rows, columns=header, dtype=object))
    validated, errors = validate_rows(frame, contract)
    if errors.ne("").any() or validated.duplicated(contract["key"]).any():
        raise AnalyticsError(f"Coorte {year} tem registros MT incompatíveis; publicação bloqueada.")
    if not validated.NU_ANO_INGRESSO.eq(year).all() or validated.NU_ANO_REFERENCIA.max() != 2024:
        raise AnalyticsError("Coorte/período diferente do arquivo.")
    analytical = trajectory(validated)
    if analytical.groupby("curso_id").size().ne(2025 - year).any():
        raise AnalyticsError("Coorte sem painel completo até 2024.")
    validated.to_csv(output, index=False)
    report = {"coorte": year, "registros_brasil": total, "uf_ausente_brasil": unknown,
              "registros_mt": len(validated), "cursos_mt": int(analytical.curso_id.nunique()),
              "modalidades_mt": {str(k): int(v) for k, v in analytical.modalidade.value_counts().items()},
              "source_sha256": sha256(source), "output_sha256": sha256(output),
              "dictionary_sha256": dictionary_hash, "parser_sha256": parser_hash,
              "status": "validated_private_not_public"}
    info_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", nargs="+", choices=["ibge", "enade", "coortes"], default=["ibge", "enade", "coortes"])
    parser.add_argument("--cohorts", nargs="+", type=int, choices=range(2016, 2021), default=list(range(2016, 2021)))
    args = parser.parse_args()
    DEST.mkdir(parents=True, exist_ok=True)
    report_path = DEST / "report.json"
    report = json.loads(report_path.read_text(encoding="utf-8")) if report_path.exists() else {}
    for name in args.only:
        if name == "ibge":
            report[name] = prepare_ibge()
        elif name == "enade":
            report[name] = prepare_enade()
        else:
            report.setdefault("coortes", {})
            for year in args.cohorts:
                report["coortes"][str(year)] = prepare_cohort(year)
                print(f"Coorte {year}: validada", flush=True)
        report["updated_at"] = datetime.now(timezone.utc).isoformat()
        report["status"] = "validated_private_not_public"
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(name + ": preparação concluída", flush=True)


if __name__ == "__main__":
    main()
